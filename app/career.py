"""Deterministic, evidence-bound helpers for the Life OS Career Agent."""

from io import BytesIO
import re


SCORE_WEIGHTS = {
    "role_fit": 25,
    "required_skills": 35,
    "preferred_skills": 10,
    "location": 10,
    "early_career_fit": 15,
    "domain_fit": 5,
}


def _normalize(value: str) -> str:
    return re.sub(r"[^a-z0-9+]+", " ", (value or "").lower()).strip()


def _contains(haystack: str, needle: str) -> bool:
    return _normalize(needle) in _normalize(haystack)


def score_job(job: dict, profile: dict) -> dict:
    """Return a transparent 0-100 score and human-readable evidence."""
    title = _normalize(job.get("title", ""))
    body = _normalize(job.get("description", ""))
    location = _normalize(f"{job.get('location', '')} {job.get('work_mode', '')}")
    profile_skills = {_normalize(skill) for skill in profile.get("skills", [])}
    required = job.get("required_skills", [])
    preferred = job.get("preferred_skills", [])

    role_match = max(
        (_contains(title, role) or _contains(role, title) for role in profile.get("target_roles", [])),
        default=False,
    )
    required_hits = [skill for skill in required if _normalize(skill) in profile_skills]
    preferred_hits = [skill for skill in preferred if _normalize(skill) in profile_skills]
    missing_required = [skill for skill in required if _normalize(skill) not in profile_skills]
    location_match = any(_contains(location, place) for place in profile.get("locations", []))
    senior_text = _normalize(f"{job.get('required_experience', '')} {title} {body}")
    senior_role = bool(re.search(r"\b(senior|staff|lead|principal|[5-9]\+? years?)\b", senior_text))
    early_career = not senior_role
    domain_terms = profile.get("project_terms", [])
    domain_hits = [term for term in domain_terms if _contains(body, term)]

    breakdown = {
        "role_fit": SCORE_WEIGHTS["role_fit"] if role_match else 0,
        "required_skills": round(SCORE_WEIGHTS["required_skills"] * len(required_hits) / max(len(required), 1)),
        "preferred_skills": round(SCORE_WEIGHTS["preferred_skills"] * len(preferred_hits) / max(len(preferred), 1)),
        "location": SCORE_WEIGHTS["location"] if location_match else 0,
        "early_career_fit": SCORE_WEIGHTS["early_career_fit"] if early_career else 0,
        "domain_fit": SCORE_WEIGHTS["domain_fit"] if domain_hits else 0,
    }
    strengths = []
    if role_match:
        strengths.append("Target role match")
    if required_hits:
        strengths.append(f"Required skills: {', '.join(required_hits)}")
    if preferred_hits:
        strengths.append(f"Preferred skills: {', '.join(preferred_hits)}")
    if location_match:
        strengths.append("Location or work mode match")
    if early_career:
        strengths.append("Early-career level")

    gaps = [f"Missing required skill: {skill}" for skill in missing_required]
    if senior_role:
        gaps.append("Role appears senior for the current profile")
    if not location_match:
        gaps.append("Location is outside the preferred search area")

    return {
        "score": min(100, sum(breakdown.values())),
        "breakdown": breakdown,
        "strengths": strengths,
        "gaps": gaps,
    }


def verify_claims(claims: list[str], supported_facts: list[str]) -> list[str]:
    """Return claims that are not exact, normalized matches to verified facts."""
    supported = {_normalize(fact) for fact in supported_facts}
    return [claim for claim in claims if _normalize(claim) not in supported]


def plan_application_fields(fields: list[dict]) -> dict:
    """Classify application fields without submitting or guessing answers."""
    safe, needs_review = [], []
    for field in fields:
        label = _normalize(field.get("label", ""))
        kind = _normalize(field.get("type", "text"))
        if "captcha" in label or kind == "captcha":
            needs_review.append({**field, "reason": "CAPTCHA requires manual completion"})
        elif field.get("required") and not field.get("value"):
            needs_review.append({**field, "reason": "Required answer is missing"})
        elif field.get("ambiguous"):
            needs_review.append({**field, "reason": "Ambiguous question requires human judgment"})
        else:
            safe.append(field)
    return {"safe_to_fill": safe, "needs_review": needs_review, "may_submit": False}


def build_application_document(application: dict, job: dict, kind: str, file_format: str) -> BytesIO:
    """Build an in-memory DOCX or PDF from verified claims only."""
    if kind not in {"cv", "motivation"} or file_format not in {"docx", "pdf"}:
        raise ValueError("Unsupported document type")
    if application.get("unsupported_claims"):
        raise ValueError("Cannot export documents with unsupported claims")

    title = "Tailored CV" if kind == "cv" else "Motivation Letter"
    lines = (
        application["cv_claims"]
        if kind == "cv"
        else application["motivation_letter"].splitlines()
    )
    target = f"{job['title']} — {job['company']}"
    stream = BytesIO()

    if file_format == "docx":
        from docx import Document
        document = Document()
        document.add_heading(title, level=0)
        document.add_paragraph(target)
        if kind == "cv":
            document.add_heading("Verified profile", level=1)
            for line in lines:
                document.add_paragraph(line, style="List Bullet")
        else:
            for line in lines:
                document.add_paragraph(line)
        document.save(stream)
    else:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        pdf = canvas.Canvas(stream, pagesize=A4)
        width, height = A4
        y = height - 62
        pdf.setFont("Helvetica-Bold", 18)
        pdf.drawString(54, y, title)
        y -= 24
        pdf.setFont("Helvetica", 10)
        pdf.drawString(54, y, target[:95])
        y -= 32
        for raw_line in lines:
            prefix = "• " if kind == "cv" else ""
            words = f"{prefix}{raw_line}".split()
            current = ""
            for word in words:
                if pdf.stringWidth(f"{current} {word}", "Helvetica", 10) > width - 108:
                    pdf.drawString(54, y, current)
                    y -= 15
                    current = word
                else:
                    current = f"{current} {word}".strip()
            if current:
                pdf.drawString(54, y, current)
            y -= 18
            if y < 54:
                pdf.showPage()
                pdf.setFont("Helvetica", 10)
                y = height - 54
        pdf.save()
    stream.seek(0)
    return stream
